import os
import shutil
import logging
import json
import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from google_auth_oauthlib.flow import InstalledAppFlow
import requests
import fitz               # PyMuPDF for PDF extraction
from pptx import Presentation  # for PowerPoint extraction
import numpy as np
from sentence_transformers import SentenceTransformer
import mysql.connector
from mysql.connector import Error
from dotenv import load_dotenv
import sys
sys.stdout.reconfigure(encoding='utf-8')
load_dotenv()  # Load .env file

# --------------------- CONFIGURATION ---------------------
USER_DB = "users.txt"  # Still using this file for file-based user registration
SETTINGS_DB = "settings.json"
LOG_FILE = "file_sorter.log"
logging.basicConfig(filename=LOG_FILE, level=logging.INFO, format='%(message)s')

# --------------------- LOAD SENTENCE-TRANSFORMERS MODEL ---------------------
model = SentenceTransformer('all-mpnet-base-v2')

# --------------------- DATABASE CONNECTION FUNCTIONS ---------------------
def create_connection():
    connection = None
    try:
        connection = mysql.connector.connect(
            host=os.getenv('DB_HOST'),        # your MySQL host,
            database=os.getenv('DB_NAME'),         # your database name
            user=os.getenv('DB_USER'),           # your MySQL username
            password=os.getenv('DB_PASSWORD')              # your MySQL password
        )
        if connection.is_connected():
            print("Successfully connected to MySQL database 'fyp'.")
            return connection
    except Error as e:
        print(f"Error connecting to MySQL database: {e}")
        return None

def close_connection(connection):
    if connection and connection.is_connected():
        connection.close()
        print("MySQL connection closed.")

# --------------------- USER AUTHENTICATION (DB) ---------------------
def login_user_db(username, password):
    """
    Validate user credentials against the database.
    Returns the user record (a dict) if successful, otherwise None.
    """
    conn = create_connection()
    if not conn:
        return None
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT * FROM users WHERE username = %s", (username,))
        user = cursor.fetchone()
        # For demonstration, we compare plaintext passwords.
        if user and user['password'] == password:
            print("Login successful!")
            return user
        else:
            print("Invalid credentials.")
            return None
    except Error as e:
        print(f"Error during login: {e}")
        return None
    finally:
        cursor.close()
        close_connection(conn)

# --------------------- RULES DATABASE FUNCTIONS ---------------------
def create_rule_db(user_id, name, source_folder, destination_folder, modifications, is_running=False):
    """
    Insert a new rule into the database.
    Returns the new rule's ID if successful, otherwise None.
    """
    conn = create_connection()
    if not conn:
        return None
    try:
        cursor = conn.cursor()
        modifications_str = json.dumps(modifications)
        query = """
            INSERT INTO rules (user_id, name, source_folder, destination_folder, modifications, is_running)
            VALUES (%s, %s, %s, %s, %s, %s)
        """
        cursor.execute(query, (user_id, name, source_folder, destination_folder, modifications_str, is_running))
        conn.commit()
        rule_id = cursor.lastrowid
        print(f"Rule created successfully with ID: {rule_id}")
        return rule_id
    except Error as e:
        print(f"Error creating rule: {e}")
        return None
    finally:
        cursor.close()
        close_connection(conn)

def get_rules_by_user_db(user_id):
    """
    Retrieve all rules for a specific user.
    Returns a list of rule dictionaries.
    """
    conn = create_connection()
    if not conn:
        return []
    try:
        cursor = conn.cursor(dictionary=True)
        query = "SELECT * FROM rules WHERE user_id = %s"
        cursor.execute(query, (user_id,))
        rules = cursor.fetchall()
        for rule in rules:
            rule["modifications"] = json.loads(rule["modifications"]) if rule["modifications"] else []
        return rules
    except Error as e:
        print(f"Error retrieving rules: {e}")
        return []
    finally:
        cursor.close()
        close_connection(conn)

def update_rule_db(rule_id, **kwargs):
    """
    Update a rule's fields in the database.
    Returns True if successful, False otherwise.
    """
    if not kwargs:
        print("No fields provided for update.")
        return False
    conn = create_connection()
    if not conn:
        return False
    try:
        cursor = conn.cursor()
        if "modifications" in kwargs:
            kwargs["modifications"] = json.dumps(kwargs["modifications"])
        fields = []
        values = []
        for key, value in kwargs.items():
            fields.append(f"{key} = %s")
            values.append(value)
        values.append(rule_id)
        query = f"UPDATE rules SET {', '.join(fields)} WHERE id = %s"
        cursor.execute(query, tuple(values))
        conn.commit()
        print(f"Rule {rule_id} updated successfully.")
        return True
    except Error as e:
        print(f"Error updating rule: {e}")
        return False
    finally:
        cursor.close()
        close_connection(conn)

def delete_rule_db(rule_id):
    """
    Delete a rule from the database.
    Returns True if successful, False otherwise.
    """
    conn = create_connection()
    if not conn:
        return False
    try:
        cursor = conn.cursor()
        query = "DELETE FROM rules WHERE id = %s"
        cursor.execute(query, (rule_id,))
        conn.commit()
        print(f"Rule {rule_id} deleted successfully.")
        return True
    except Error as e:
        print(f"Error deleting rule: {e}")
        return False
    finally:
        cursor.close()
        close_connection(conn)

# --------------------- HELPER FUNCTIONS (FILE SORTING) ---------------------
def extract_text_from_pdf(pdf_path):
    """Extracts text from a PDF file using PyMuPDF."""
    text = ""
    try:
        doc = fitz.open(pdf_path)
        for page in doc:
            text += page.get_text()
        doc.close()
    except Exception as e:
        print(f"Error extracting text from {pdf_path}: {e}")
    return text

def extract_text_from_pptx(pptx_path):
    """Extracts text from a PowerPoint file using python-pptx."""
    text = ""
    try:
        prs = Presentation(pptx_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
    except Exception as e:
        print(f"Error extracting text from {pptx_path}: {e}")
    return text

def cosine_similarity(vec1, vec2):
    """Computes cosine similarity between two vectors."""
    return np.dot(vec1, vec2) / (np.linalg.norm(vec1) * np.linalg.norm(vec2) + 1e-8)

def classify_file_content_local(content, category_examples):
    """
    Uses a local embedding model to classify file content.
    Returns the category with the highest cosine similarity.
    """
    try:
        content_embedding = model.encode(content)
    except Exception as e:
        print(f"Error encoding content: {e}")
        return None
    best_category = None
    best_similarity = -1
    for category, example_text in category_examples.items():
        try:
            example_embedding = model.encode(example_text)
            sim = cosine_similarity(content_embedding, example_embedding)
            print(f"Similarity for {category}: {sim}")
            if sim > best_similarity:
                best_similarity = sim
                best_category = category
        except Exception as e:
            print(f"Error encoding example for {category}: {e}")
    return best_category

def sort_files_by_content_with_ai(source_folder, categories, destination_map):
    """
    Sorts files from source_folder based on file content using local AI.
    """
    category_examples = {
        'ec2324': "Course EC2324 covers artificial intelligence techniques, including machine learning, neural networks, and NLP.",
        'ec2306': "Course EC2306 covers Java programming and OOP, including classes, inheritance, design patterns.",
        'ec2312': "Course EC2312 covers OS concepts such as processes, Linux, concurrency, file systems.",
        'ec2302': "Course EC2302 covers software engineering, SDLC, design patterns, testing, project management.",
        'ec2320': "Course EC2320 covers final year project, summative, supervisor, coordinator.",
        'non-related': "This document is not related to any of the defined courses."
    }
    for filename in os.listdir(source_folder):
        file_path = os.path.join(source_folder, filename)
        if os.path.isfile(file_path):
            content = ""
            if filename.lower().endswith(".pdf"):
                content = extract_text_from_pdf(file_path)
                print(f"Extracted text from {filename}: {content[:100].encode('ascii', 'ignore').decode()}")
            elif filename.lower().endswith(".pptx"):
                content = extract_text_from_pptx(file_path)
                print(f"Extracted text from {filename}: {content[:100].encode('ascii', 'ignore').decode()}")
            else:
                try:
                    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                        f.write(f"Extracted text from {filename}: {content[:100]}\n")
                        content = f.read(1024)
                except Exception as e:
                    print(f"Could not read file {filename}: {e}")
                    continue
            category = classify_file_content_local(content, category_examples)
            print(f"File {filename} classified as: '{category}'")
            if category and category in destination_map:
                dest_folder = destination_map[category]
                os.makedirs(dest_folder, exist_ok=True)
                dest_path = os.path.join(dest_folder, filename)
                try:
                    shutil.move(file_path, dest_path)
                    current_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    logging.info(f"{current_time} | {category} | Moved to {dest_folder} | {filename}")
                    print(f"Moved {filename} to {dest_folder}")
                except Exception as e:
                    print(f"Error moving {filename}: {e}")
            else:
                print(f"File {filename} could not be classified or category not found.")
    
def sort_files_by_date(source_folder, destination_base):
    for filename in os.listdir(source_folder):
        file_path = os.path.join(source_folder, filename)
        if os.path.isfile(file_path):
            try:
                mod_time = os.path.getmtime(file_path)
                date_str = datetime.datetime.fromtimestamp(mod_time).strftime("%Y-%m-%d")
                dest_folder = os.path.join(destination_base, date_str)
                os.makedirs(dest_folder, exist_ok=True)
                shutil.move(file_path, os.path.join(dest_folder, filename))
                current_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                logging.info(f"{current_time} | DateSort | Moved to {dest_folder} | {filename}")
                print(f"Moved {filename} to {dest_folder}")
            except Exception as e:
                print(f"Error sorting {filename} by date: {e}")

def sort_files_by_file_type(source_folder, destination_base):
    for filename in os.listdir(source_folder):
        file_path = os.path.join(source_folder, filename)
        if os.path.isfile(file_path):
            ext = os.path.splitext(filename)[1].lower()
            ext = ext.lstrip('.') if ext else "unknown"
            dest_folder = os.path.join(destination_base, ext)
            os.makedirs(dest_folder, exist_ok=True)
            try:
                shutil.move(file_path, os.path.join(dest_folder, filename))
                current_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                logging.info(f"{current_time} | FileTypeSort | Moved to {dest_folder} | {filename}")
                print(f"Moved {filename} to {dest_folder}")
            except Exception as e:
                print(f"Error sorting {filename} by file type: {e}")

def sort_files_by_file_name(source_folder, destination_base):
    for filename in os.listdir(source_folder):
        file_path = os.path.join(source_folder, filename)
        if os.path.isfile(file_path):
            first_letter = filename[0].upper() if filename else "Unknown"
            dest_folder = os.path.join(destination_base, first_letter)
            os.makedirs(dest_folder, exist_ok=True)
            try:
                shutil.move(file_path, os.path.join(dest_folder, filename))
                current_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                logging.info(f"{current_time} | FileNameSort | Moved to {dest_folder} | {filename}")
                print(f"Moved {filename} to {dest_folder}")
            except Exception as e:
                print(f"Error sorting {filename} by file name: {e}")

# --------------------- MAIN APPLICATION (GUI) ---------------------
class FileSorterApp:
    def __init__(self, root):
        self.root = root
        
        # Initialize default settings
        self.settings = {
            "ai_model": "all-mpnet-base-v2",
            "default_sort_folder": "",
            "sorting_frequency": "Manual",
            "theme": "Light",
            "font_size": 10
        }
        self.load_settings()
        
        self.root.title("Automatic File Sorter")
        self.root.geometry("900x600")
        self.root.resizable(True, True)
        
        # Configure ttk styles
        self.style = ttk.Style()
        self.style.theme_use("clam")
        self.style.configure('TFrame', padding=10)
        self.style.configure('TLabel', padding=5)
        self.style.configure('TButton', padding=5)
        
        # Top header frame with settings button
        header_frame = ttk.Frame(self.root)
        header_frame.pack(fill='x')
        header_label = ttk.Label(header_frame, text="Automatic File Sorter", font=("Helvetica", 16))
        header_label.pack(side=tk.LEFT, padx=10, pady=10)
        settings_button = ttk.Button(header_frame, text="⚙️", width=3, command=self.open_settings_popup)
        settings_button.pack(side=tk.RIGHT, padx=10, pady=10)
        
        # Notebook with tabs: Login, Logs, Rules
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(expand=True, fill='both')
        self.login_frame = ttk.Frame(self.notebook)
        self.logs_frame = ttk.Frame(self.notebook)
        self.rules_frame = ttk.Frame(self.notebook)
        self.notebook.add(self.login_frame, text="Login")
        self.notebook.add(self.logs_frame, text="Logs")
        self.notebook.add(self.rules_frame, text="Rules")
        
        self.build_login_tab()
        self.build_logs_tab()
        
        # Disable Logs and Rules tabs until a user logs in
        self.notebook.tab(1, state='disabled')
        self.notebook.tab(2, state='disabled')
        
        # Rules data will now come from the DB
        self.rules_data = []
        self.copied_rule = None
        
        # Advanced monitor settings variables
        self.monitor_interval = tk.IntVar(value=10)
        self.monitor_unit = tk.StringVar(value="minutes")
        self.time_sort_ascending = True
        self.all_log_entries = []
        self.new_rule_frame = None
        
        # Always available source/destination paths
        self.source_path = tk.StringVar(value="Downloads")
        self.dest_path = tk.StringVar(value=r"C:\Users\Sam\Documents\test")
        
        # Current logged in user (as a dict from the DB)
        self.current_user = None

    # --------------------- SETTINGS ---------------------
    def load_settings(self):
        if os.path.exists(SETTINGS_DB):
            with open(SETTINGS_DB, "r") as file:
                self.settings = json.load(file)
    
    def save_settings(self):
        self.settings["ai_model"] = self.ai_model_var.get()
        self.settings["default_sort_folder"] = self.default_sort_folder.get()
        self.settings["sorting_frequency"] = self.sorting_frequency_var.get()
        self.settings["theme"] = self.theme_var.get()
        self.settings["font_size"] = self.font_size_var.get()
        with open(SETTINGS_DB, "w") as file:
            json.dump(self.settings, file, indent=4)
        global model
        model = SentenceTransformer(self.ai_model_var.get())
        self.apply_theme()
        messagebox.showinfo("Settings", "Settings saved successfully!")
    
    def export_settings(self):
        export_path = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("JSON Files", "*.json")])
        if export_path:
            with open(export_path, "w") as file:
                json.dump(self.settings, file, indent=4)
            messagebox.showinfo("Export Settings", "Settings exported successfully!")
    
    def import_settings(self):
        import_path = filedialog.askopenfilename(filetypes=[("JSON Files", "*.json")])
        if import_path:
            with open(import_path, "r") as file:
                self.settings = json.load(file)
            self.ai_model_var.set(self.settings.get("ai_model", "all-mpnet-base-v2"))
            self.default_sort_folder.set(self.settings.get("default_sort_folder", ""))
            self.sorting_frequency_var.set(self.settings.get("sorting_frequency", "Manual"))
            self.theme_var.set(self.settings.get("theme", "Light"))
            self.font_size_var.set(self.settings.get("font_size", 10))
            self.save_settings()
            messagebox.showinfo("Import Settings", "Settings imported successfully!")
    
    # --------------------- THEME APPLIER ---------------------
    def apply_theme(self):
        font_size = self.font_size_var.get()
        if self.theme_var.get() == "Dark":
            self.root.configure(bg="gray20")
            self.style.configure("TFrame", background="gray20")
            self.style.configure("TLabel", background="gray20", foreground="white", font=("Helvetica", font_size))
            self.style.configure("TButton", background="gray30", foreground="white", font=("Helvetica", font_size))
        else:
            self.root.configure(bg="SystemButtonFace")
            self.style.configure("TFrame", background="SystemButtonFace")
            self.style.configure("TLabel", background="SystemButtonFace", foreground="black", font=("Helvetica", font_size))
            self.style.configure("TButton", background="SystemButtonFace", foreground="black", font=("Helvetica", font_size))
    
    # --------------------- RULES DATA (FROM DB) ---------------------
    def load_rules_from_db(self):
        """Loads rules for the current user from the database."""
        if self.current_user:
            self.rules_data = get_rules_by_user_db(self.current_user["id"])
            print("Loaded rules from DB:", self.rules_data)
        else:
            self.rules_data = []
            print("No user logged in – rules list is empty")
    
    # --------------------- RULE ACTIONS (DB FUNCTIONS) ---------------------
    def toggle_rule_running(self, rule):
        new_state = not rule["is_running"]
        if update_rule_db(rule["id"], is_running=new_state):
            rule["is_running"] = new_state
            if rule["is_running"]:
                self.continuous_sorting(rule)
                messagebox.showinfo("Rule Activated", f"Rule '{rule['name']}' is now running continuously.")
            else:
                messagebox.showinfo("Rule Stopped", f"Rule '{rule['name']}' has been stopped.")
            self.load_rules_from_db()
            self.build_rules_tab_ui()
        else:
            messagebox.showerror("Error", "Failed to update rule state.")
    
    def continuous_sorting(self, rule):
        if not rule.get("is_running"):
            return
        source_folder = rule.get("source_folder", self.source_path.get())
        destination_base = rule.get("destination_folder", self.dest_path.get())
        mods = rule.get("modifications", [])
        if "ai_file_content" in mods:
            categories = ['ec2324', 'ec2306', 'ec2312', 'ec2302', 'ec2320' '2170', 'non-related']
            destination_map = {
                'ec2324': os.path.join(source_folder, "Sorted_ec2324"),
                'ec2306': os.path.join(source_folder, "Sorted_ec2306"),
                'ec2312': os.path.join(source_folder, "Sorted_ec2312"),
                'ec2302': os.path.join(source_folder, "Sorted_ec2302"),
                'ec2320': os.path.join(source_folder, "Sorted_ec2320"),
                '2170': os.path.join(source_folder, "Sorted_2170"),
                'non-related': os.path.join(source_folder, "Sorted_non-related")
            }
            sort_files_by_content_with_ai(source_folder, categories, destination_map)
        if "date" in mods:
            sort_files_by_date(source_folder, destination_base)
        if "file_type" in mods:
            sort_files_by_file_type(source_folder, destination_base)
        if "file_name" in mods:
            sort_files_by_file_name(source_folder, destination_base)
        self.root.after(60000, lambda: self.continuous_sorting(rule))
    
    def delete_rule(self, rule):
        confirm = messagebox.askyesno("Delete", f"Are you sure you want to delete rule '{rule['name']}'?")
        if confirm:
            if delete_rule_db(rule["id"]):
                messagebox.showinfo("Delete", "Rule deleted successfully.")
            else:
                messagebox.showerror("Error", "Failed to delete rule.")
            self.load_rules_from_db()
            self.build_rules_tab_ui()
    
    # --------------------- RULES TAB UI (DB) ---------------------
    def build_rules_tab_ui(self):
        for child in self.rules_frame.winfo_children():
            child.destroy()
        
        top_bar = ttk.Frame(self.rules_frame)
        top_bar.pack(fill='x', pady=10, padx=10)
        add_rule_btn = ttk.Button(top_bar, text="Add Rule", command=self.create_new_rule_tab)
        add_rule_btn.pack(side=tk.LEFT, padx=(0,10))
        paste_btn = ttk.Button(top_bar, text="Paste", command=self.paste_rule)
        paste_btn.pack(side=tk.LEFT)
        
        self.rules_container = ttk.Frame(self.rules_frame)
        self.rules_container.pack(fill='both', expand=True, padx=10, pady=10)
        
        for rule in self.rules_data:
            self.build_rule_row(rule)
    
    def build_rule_row(self, rule):
        row_frame = ttk.Frame(self.rules_container, padding=5)
        row_frame.pack(fill='x', pady=5)
        
        left_frame = ttk.Frame(row_frame)
        left_frame.pack(side=tk.LEFT, anchor=tk.W, expand=True, fill='x')
        
        if rule.get("is_running"):
            start_pause_text = "⏸️"
            status_text = "Running"
        else:
            start_pause_text = "▶️"
            status_text = "Stopped"
        
        start_pause_btn = ttk.Button(left_frame, text=start_pause_text, width=3,
                                     command=lambda r=rule: self.toggle_rule_running(r))
        start_pause_btn.pack(side=tk.LEFT, padx=(0,10))
        
        info_frame = ttk.Frame(left_frame)
        info_frame.pack(side=tk.LEFT, fill='x', expand=True)
        lbl_name = ttk.Label(info_frame, text=rule.get("name", "Unnamed"), font=("Helvetica", 11, "bold"))
        lbl_name.pack(anchor=tk.W)
        lbl_status = ttk.Label(info_frame, text=status_text,
                               foreground="green" if rule.get("is_running") else "gray")
        lbl_status.pack(anchor=tk.W)
        
        right_frame = ttk.Frame(row_frame)
        right_frame.pack(side=tk.RIGHT)
        log_btn = ttk.Button(right_frame, text="Log", command=self.go_to_logs)
        log_btn.pack(side=tk.LEFT, padx=(0,5))
        edit_btn = ttk.Button(right_frame, text="Edit", command=lambda r=rule: self.edit_rule(r))
        edit_btn.pack(side=tk.LEFT, padx=(0,5))
        menu_btn = ttk.Menubutton(right_frame, text="⋮")
        menu = tk.Menu(menu_btn, tearoff=0)
        menu.add_command(label="Delete", command=lambda r=rule: self.delete_rule(r))
        menu.add_command(label="Copy", command=lambda r=rule: self.copy_rule(r))
        menu_btn["menu"] = menu
        menu_btn.pack(side=tk.LEFT)
    
    def paste_rule(self):
        if not self.copied_rule:
            messagebox.showinfo("Paste", "No rule has been copied yet.")
            return
        new_rule = dict(self.copied_rule)
        new_rule["name"] = new_rule["name"] + " (Copy)"
        new_rule["is_running"] = False
        rule_id = create_rule_db(self.current_user["id"], new_rule["name"], self.source_path.get(),
                                  self.dest_path.get(), new_rule.get("modifications", []), is_running=False)
        if rule_id:
            messagebox.showinfo("Success", "Rule copied successfully.")
        else:
            messagebox.showerror("Error", "Failed to copy rule.")
        self.load_rules_from_db()
        self.build_rules_tab_ui()
    
    def edit_rule(self, rule):
        self.create_new_rule_tab(edit_rule=rule)
    
    def copy_rule(self, rule):
        self.copied_rule = dict(rule)
        messagebox.showinfo("Copy", f"Copied rule: {rule['name']}")
    
    # --------------------- NEW RULE TAB ---------------------
    def create_new_rule_tab(self, edit_rule=None):
        if self.new_rule_frame:
            self.notebook.select(self.new_rule_frame)
            return
        self.new_rule_frame = ttk.Frame(self.notebook)
        self.notebook.add(self.new_rule_frame, text="New Rule")
        self.notebook.select(self.new_rule_frame)
        container = ttk.Frame(self.new_rule_frame, padding=10)
        container.pack(expand=True, fill='both')
        
        # Row 1: Name
        row1 = ttk.Frame(container)
        row1.pack(fill='x', pady=5)
        ttk.Label(row1, text="Name", width=15).pack(side=tk.LEFT, anchor=tk.NW, padx=(0,5))
        self.rule_name = tk.StringVar(value=edit_rule["name"] if edit_rule else "")
        ttk.Entry(row1, textvariable=self.rule_name, width=50).pack(side=tk.LEFT, padx=(0,20))
        
        # Row 2: Monitor
        row2 = ttk.Frame(container)
        row2.pack(fill='x', pady=10)
        ttk.Label(row2, text="Monitor", width=15).pack(side=tk.LEFT, anchor=tk.NW, padx=(0,5))
        self.source_path = tk.StringVar(value=edit_rule["source_folder"] if edit_rule else "Downloads")
        ttk.Entry(row2, textvariable=self.source_path, width=40, state='readonly').pack(side=tk.LEFT, padx=(0,5))
        ttk.Button(row2, text="📂", command=self.choose_source_directory, width=3).pack(side=tk.LEFT, padx=(0,5))
        
        # Row 3: Modification
        row3 = ttk.Frame(container)
        row3.pack(fill='x', pady=10)
        ttk.Label(row3, text="Modification", width=15).pack(side=tk.LEFT, anchor=tk.NW, padx=(0,5))
        mod_picker_frame = ttk.Frame(row3)
        mod_picker_frame.pack(side=tk.LEFT, padx=(0,10))
        self.mod_options = ["date", "file_name", "file_type", "ai_file_content"]
        self.selected_mod = tk.StringVar(value=self.mod_options[0])
        ttk.Combobox(mod_picker_frame, textvariable=self.selected_mod, values=self.mod_options, width=15).pack(side=tk.LEFT, padx=(0,5))
        ttk.Button(mod_picker_frame, text="Add", command=self.add_modification).pack(side=tk.LEFT)
        self.mods_frame = ttk.Frame(row3)
        self.mods_frame.pack(side=tk.LEFT, fill='x')
        self.modifications = edit_rule["modifications"] if edit_rule else []
        self.refresh_mod_chips()
        
        # Row 4: Destination
        row4 = ttk.Frame(container)
        row4.pack(fill='x', pady=10)
        ttk.Label(row4, text="Then", width=15).pack(side=tk.LEFT, anchor=tk.NW, padx=(0,5))
        ttk.Label(row4, text="Move files to").pack(side=tk.LEFT, padx=(0,5))
        self.dest_path = tk.StringVar(value=edit_rule["destination_folder"] if edit_rule else r"C:\Users\Sam\Documents\dcs\fyp")
        ttk.Entry(row4, textvariable=self.dest_path, width=50, state='readonly').pack(side=tk.LEFT, padx=(5,5))
        ttk.Button(row4, text="📂", command=self.choose_dest_directory, width=3).pack(side=tk.LEFT)
        
        # Save Rule button
        ttk.Button(container, text="Save Rule", command=lambda: self.save_new_rule(edit_rule)).pack(pady=10)
    
    def save_new_rule(self, edit_rule=None):
        rule_name = self.rule_name.get() or "Unnamed Rule"
        user_id = self.current_user["id"] if self.current_user else None
        if not user_id:
            messagebox.showerror("Error", "You must be logged in to save a rule.")
            return
        source_folder = self.source_path.get()
        destination_folder = self.dest_path.get()
        modifications = self.modifications
        
        if edit_rule:
            if update_rule_db(edit_rule["id"], name=rule_name, source_folder=source_folder,
                              destination_folder=destination_folder, modifications=modifications):
                messagebox.showinfo("Success", f"Rule '{rule_name}' updated successfully.")
            else:
                messagebox.showerror("Error", "Failed to update rule.")
        else:
            rule_id = create_rule_db(user_id, rule_name, source_folder, destination_folder, modifications, is_running=False)
            if rule_id:
                messagebox.showinfo("Success", f"New rule '{rule_name}' saved!")
            else:
                messagebox.showerror("Error", "Failed to save new rule.")
        
        self.load_rules_from_db()
        self.build_rules_tab_ui()
        if self.new_rule_frame:
            self.notebook.forget(self.new_rule_frame)
            self.new_rule_frame = None
    
    def choose_source_directory(self):
        folder_selected = filedialog.askdirectory()
        if folder_selected:
            self.source_path.set(folder_selected)
    
    def choose_dest_directory(self):
        folder_selected = filedialog.askdirectory()
        if folder_selected:
            self.dest_path.set(folder_selected)
    
    def open_monitor_settings(self):
        monitor_win = tk.Toplevel(self.root)
        monitor_win.title("Monitor Settings")
        monitor_win.geometry("300x150")
        monitor_win.resizable(False, False)
        ttk.Label(monitor_win, text="Examine folder every:").pack(pady=10)
        row = ttk.Frame(monitor_win)
        row.pack(pady=5)
        ttk.Spinbox(row, from_=1, to=1440, textvariable=self.monitor_interval, width=5).pack(side=tk.LEFT, padx=(0,5))
        ttk.Combobox(row, textvariable=self.monitor_unit, values=["minutes", "hours", "days"], width=8).pack(side=tk.LEFT)
        ttk.Button(monitor_win, text="Save", command=monitor_win.destroy).pack(pady=10)
    
    def add_modification(self):
        mod = self.selected_mod.get()
        if mod and mod not in self.modifications:
            self.modifications.append(mod)
            self.refresh_mod_chips()
        else:
            messagebox.showinfo("Info", f"'{mod}' is already added or invalid.")
    
    def remove_modification(self, mod):
        if mod in self.modifications:
            self.modifications.remove(mod)
            self.refresh_mod_chips()
    
    def refresh_mod_chips(self):
        for child in self.mods_frame.winfo_children():
            child.destroy()
        for mod in self.modifications:
            chip_frame = ttk.Frame(self.mods_frame, padding=5)
            chip_frame.pack(side=tk.LEFT, padx=5)
            ttk.Label(chip_frame, text=mod).pack(side=tk.LEFT)
            ttk.Button(chip_frame, text="x", width=2, command=lambda m=mod: self.remove_modification(m)).pack(side=tk.LEFT)
    
    def open_settings_popup(self):
        settings_win = tk.Toplevel(self.root)
        settings_win.title("Settings")
        settings_win.geometry("600x500")
        settings_win.resizable(True, True)
        settings_notebook = ttk.Notebook(settings_win)
        settings_notebook.pack(expand=True, fill='both', padx=10, pady=10)
        
        ai_sort_frame = ttk.Frame(settings_notebook)
        settings_notebook.add(ai_sort_frame, text="AI & Sorting")
        self.build_ai_sort_settings(ai_sort_frame)
        
        account_ui_frame = ttk.Frame(settings_notebook)
        settings_notebook.add(account_ui_frame, text="Account & UI")
        self.build_account_ui_settings(account_ui_frame)
        
        backup_frame = ttk.Frame(settings_notebook)
        settings_notebook.add(backup_frame, text="Backup & Data")
        self.build_backup_settings(backup_frame)
        
        ttk.Button(settings_win, text="Save Settings", command=self.save_settings).pack(pady=10)
    
    def build_default_sort_folder_section(self, parent):
        ttk.Label(parent, text="Default Sorting Folder:").grid(row=0, column=0, sticky=tk.W, padx=5, pady=5)
        self.default_sort_folder = tk.StringVar(value=self.settings.get("default_sort_folder", ""))
        self.sort_folder_entry = ttk.Entry(parent, textvariable=self.default_sort_folder, width=40, state='readonly')
        self.sort_folder_entry.grid(row=0, column=1, padx=5, pady=5)
        ttk.Button(parent, text="Browse", command=self.choose_default_sort_folder).grid(row=0, column=2, padx=5, pady=5)
    
    def build_ai_sort_settings(self, parent):
        ai_frame = ttk.LabelFrame(parent, text="AI Model Selection & Configuration", padding=10)
        ai_frame.pack(fill='x', padx=5, pady=5)
        ttk.Label(ai_frame, text="Choose AI Model:").grid(row=0, column=0, sticky=tk.W, padx=5, pady=5)
        self.ai_model_var = tk.StringVar(value=self.settings.get("ai_model", "all-mpnet-base-v2"))
        self.ai_model_combo = ttk.Combobox(
            ai_frame, 
            textvariable=self.ai_model_var, 
            values=["all-mpnet-base-v2", "all-MiniLM-L6-v2", "paraphrase-distilroberta-base-v1"]
        )
        self.ai_model_combo.grid(row=0, column=1, padx=5, pady=5)
        
        sort_frame = ttk.LabelFrame(parent, text="Sorting & Organization Preferences", padding=10)
        sort_frame.pack(fill='x', padx=5, pady=5)
        self.build_default_sort_folder_section(sort_frame)
        ttk.Label(sort_frame, text="Sorting Frequency:").grid(row=1, column=0, sticky=tk.W, padx=5, pady=5)
        self.sorting_frequency_var = tk.StringVar(value=self.settings.get("sorting_frequency", "Manual"))
        self.sort_freq_combo = ttk.Combobox(sort_frame, textvariable=self.sorting_frequency_var, values=["Real-time", "Hourly", "Daily", "Manual"])
        self.sort_freq_combo.grid(row=1, column=1, padx=5, pady=5)
    
    def build_account_ui_settings(self, parent):
        account_frame = ttk.LabelFrame(parent, text="Account & Authentication", padding=10)
        account_frame.pack(fill='x', padx=5, pady=5)
        ttk.Button(account_frame, text="Logout", command=self.logout).grid(row=0, column=0, padx=5, pady=5)
        ttk.Button(account_frame, text="Change Password", command=self.change_password).grid(row=0, column=1, padx=5, pady=5)
        
        ui_frame = ttk.LabelFrame(parent, text="UI & Appearance Settings", padding=10)
        ui_frame.pack(fill='x', padx=5, pady=5)
        ttk.Label(ui_frame, text="Theme:").grid(row=0, column=0, sticky=tk.W, padx=5, pady=5)
        self.theme_var = tk.StringVar(value=self.settings.get("theme", "Light"))
        self.theme_combo = ttk.Combobox(ui_frame, textvariable=self.theme_var, values=["Light", "Dark"])
        self.theme_combo.grid(row=0, column=1, padx=5, pady=5)
        ttk.Label(ui_frame, text="Font Size:").grid(row=1, column=0, sticky=tk.W, padx=5, pady=5)
        self.font_size_var = tk.IntVar(value=self.settings.get("font_size", 10))
        self.font_size_spin = ttk.Spinbox(ui_frame, from_=8, to=20, textvariable=self.font_size_var, width=5)
        self.font_size_spin.grid(row=1, column=1, padx=5, pady=5)
    
    def build_backup_settings(self, parent):
        backup_frame = ttk.LabelFrame(parent, text="Backup & Data Management", padding=10)
        backup_frame.pack(fill='x', padx=5, pady=5)
        ttk.Button(backup_frame, text="Export Settings", command=self.export_settings).grid(row=0, column=0, padx=5, pady=5)
        ttk.Button(backup_frame, text="Import Settings", command=self.import_settings).grid(row=0, column=1, padx=5, pady=5)
        ttk.Button(backup_frame, text="Delete All Sorted Files", command=self.delete_sorted_files).grid(row=1, column=0, columnspan=2, padx=5, pady=5)
    
    def choose_default_sort_folder(self):
        folder_selected = filedialog.askdirectory()
        if folder_selected:
            self.default_sort_folder.set(folder_selected)
    
    # --------------------- ACCOUNT ACTIONS ---------------------
    def logout(self):
        self.current_user = None
        self.notebook.tab(1, state='disabled')
        self.notebook.tab(2, state='disabled')
        self.notebook.select(0)
        messagebox.showinfo("Logout", "You have been logged out.")
    
    def change_password(self):
        if not self.current_user:
            messagebox.showerror("Error", "No user is currently logged in.")
            return
        cp_win = tk.Toplevel(self.root)
        cp_win.title("Change Password")
        cp_win.geometry("300x200")
        ttk.Label(cp_win, text="Current Password:").pack(pady=5)
        current_pw_var = tk.StringVar()
        ttk.Entry(cp_win, textvariable=current_pw_var, show="*").pack(pady=5)
        ttk.Label(cp_win, text="New Password:").pack(pady=5)
        new_pw_var = tk.StringVar()
        ttk.Entry(cp_win, textvariable=new_pw_var, show="*").pack(pady=5)
        ttk.Label(cp_win, text="Confirm New Password:").pack(pady=5)
        confirm_pw_var = tk.StringVar()
        ttk.Entry(cp_win, textvariable=confirm_pw_var, show="*").pack(pady=5)
        
        def update_password():
            current_pw = current_pw_var.get().strip()
            new_pw = new_pw_var.get().strip()
            confirm_pw = confirm_pw_var.get().strip()
            if not current_pw or not new_pw or not confirm_pw:
                messagebox.showerror("Error", "All fields are required.")
                return
            if new_pw != confirm_pw:
                messagebox.showerror("Error", "New passwords do not match.")
                return
            if os.path.exists(USER_DB):
                with open(USER_DB, "r") as file:
                    users = dict(line.strip().split(":") for line in file if ":" in line)
            else:
                messagebox.showerror("Error", "User database not found.")
                cp_win.destroy()
                return
            if self.current_user["username"] not in users or users[self.current_user["username"]] != current_pw:
                messagebox.showerror("Error", "Current password is incorrect.")
                return
            users[self.current_user["username"]] = new_pw
            with open(USER_DB, "w") as file:
                for user, pw in users.items():
                    file.write(f"{user}:{pw}\n")
            messagebox.showinfo("Success", "Password changed successfully!")
            cp_win.destroy()
        
        ttk.Button(cp_win, text="Update Password", command=update_password).pack(pady=10)
    
    # --------------------- LOGS TAB ---------------------
    def build_logs_tab(self):
        top_frame = ttk.Frame(self.logs_frame)
        top_frame.pack(fill='x', padx=10, pady=10)
        self.search_var = tk.StringVar()
        search_entry = ttk.Entry(top_frame, textvariable=self.search_var, width=30)
        search_entry.pack(side=tk.LEFT, padx=(0,5))
        search_btn = ttk.Button(top_frame, text="Search", command=self.search_logs)
        search_btn.pack(side=tk.LEFT)
        refresh_btn = ttk.Button(top_frame, text="Refresh", command=self.load_logs_into_tree)
        refresh_btn.pack(side=tk.LEFT, padx=(5,0))
        columns = ("time", "rule", "message", "file")
        self.log_tree = ttk.Treeview(self.logs_frame, columns=columns, show='headings', height=15)
        self.log_tree.heading("time", text="Time ▲", command=lambda: self.sort_by("time"))
        self.log_tree.heading("rule", text="Rule")
        self.log_tree.heading("message", text="Message")
        self.log_tree.heading("file", text="File")
        self.log_tree.column("time", width=140, anchor=tk.W)
        self.log_tree.column("rule", width=180, anchor=tk.W)
        self.log_tree.column("message", width=220, anchor=tk.W)
        self.log_tree.column("file", width=200, anchor=tk.W)
        self.log_tree.pack(expand=True, fill='both', padx=10, pady=5)
        self.all_log_entries = []
        self.load_logs_into_tree()
    
    def load_logs_into_tree(self):
        for row in self.log_tree.get_children():
            self.log_tree.delete(row)
        self.all_log_entries.clear()
        if os.path.exists(LOG_FILE):
            with open(LOG_FILE, "r") as log_file:
                for line in log_file:
                    line = line.strip()
                    if not line or "dummy" in line.lower():
                        continue
                    parts = line.split(" | ")
                    if len(parts) == 4:
                        time_str, rule_str, message_str, file_str = parts
                    else:
                        time_str = line[:19]
                        rule_str = "Unknown"
                        message_str = line[22:]
                        file_str = ""
                    entry = (time_str, rule_str, message_str, file_str)
                    self.all_log_entries.append(entry)
        for entry in self.all_log_entries:
            self.log_tree.insert("", tk.END, values=entry)
    
    def search_logs(self):
        search_term = self.search_var.get().lower().strip()
        for row in self.log_tree.get_children():
            self.log_tree.delete(row)
        for entry in self.all_log_entries:
            if search_term in " ".join(entry).lower():
                self.log_tree.insert("", tk.END, values=entry)
    
    def sort_by(self, col):
        data = [(self.log_tree.set(child, col), child) for child in self.log_tree.get_children('')]
        data.sort(reverse=not self.time_sort_ascending)
        for index, (val, child) in enumerate(data):
            self.log_tree.move(child, '', index)
        self.time_sort_ascending = not self.time_sort_ascending
        arrow = "▲" if self.time_sort_ascending else "▼"
        self.log_tree.heading("time", text=f"Time {arrow}", command=lambda: self.sort_by("time"))
    
    def go_to_logs(self):
        self.notebook.select(self.logs_frame)
    
    # --------------------- LOGIN TAB ---------------------
    def build_login_tab(self):
        self.login_page = ttk.Frame(self.login_frame)
        self.register_page = ttk.Frame(self.login_frame)
        self.login_page.place(relx=0.5, rely=0.5, anchor=tk.CENTER)
        self.register_page.place(relx=0.5, rely=0.5, anchor=tk.CENTER)
        self.build_login_page_ui()
        self.build_register_page_ui()
        self.login_page.tkraise()
    
    def build_login_page_ui(self):
        container = ttk.Frame(self.login_page, padding=30)
        container.pack()
        title_label = ttk.Label(container, text="Sign in", font=("Helvetica", 20, "bold"))
        title_label.pack(pady=(0, 5))
        subtitle_label = ttk.Label(container, text="Automatic File Sorter", font=("Helvetica", 11), foreground="#555")
        subtitle_label.pack(pady=(0, 20))
        email_label = ttk.Label(container, text="Email or Phone", font=("Helvetica", 10))
        email_label.pack(anchor="w")
        self.login_username_var = tk.StringVar()
        self.login_username_entry = ttk.Entry(container, textvariable=self.login_username_var, width=35, font=("Helvetica", 10))
        self.login_username_entry.pack(pady=(0, 15))
        password_label = ttk.Label(container, text="Password", font=("Helvetica", 10))
        password_label.pack(anchor="w")
        self.login_password_var = tk.StringVar()
        self.show_password = False
        self.login_password_entry = ttk.Entry(container, textvariable=self.login_password_var, width=35, font=("Helvetica", 10), show="*")
        self.login_password_entry.pack(side=tk.TOP, pady=(0,5))
        show_label = ttk.Label(container, text="show", foreground="blue", cursor="hand2")
        show_label.pack(anchor="w")
        show_label.bind("<Button-1>", lambda e: self.toggle_password_visibility(show_label))
        sign_in_btn = ttk.Button(container, text="Sign in", command=self.login, width=35)
        sign_in_btn.pack(pady=(0, 10))
        or_label = ttk.Label(container, text="or", foreground="#999")
        or_label.pack(pady=(10, 5))
        google_btn = ttk.Button(container, text="Sign in with Google", command=self.google_sign_in, width=35)
        google_btn.pack(pady=(5, 15))
        bottom_frame = ttk.Frame(container)
        bottom_frame.pack()
        new_label = ttk.Label(bottom_frame, text="New to AFS?", font=("Helvetica", 9))
        new_label.pack(side="left")
        join_link = ttk.Label(bottom_frame, text="Join now", foreground="blue", cursor="hand2", font=("Helvetica", 9, "underline"))
        join_link.pack(side="left", padx=(5,0))
        join_link.bind("<Button-1>", lambda e: self.show_register_page())
    
    def toggle_password_visibility(self, show_label):
        self.show_password = not self.show_password
        if self.show_password:
            self.login_password_entry.config(show="")
            show_label.config(text="hide")
        else:
            self.login_password_entry.config(show="*")
            show_label.config(text="show")
    
    def build_register_page_ui(self):
        container = ttk.Frame(self.register_page, padding=30)
        container.pack()
        title = ttk.Label(container, text="Create an Account", font=("Helvetica", 20, "bold"))
        title.pack(pady=(0, 20))
        username_label = ttk.Label(container, text="Username", font=("Helvetica", 10))
        username_label.pack(anchor="w")
        self.register_username_var = tk.StringVar()
        username_entry = ttk.Entry(container, textvariable=self.register_username_var, width=35, font=("Helvetica", 10))
        username_entry.pack(pady=(0,10))
        email_label = ttk.Label(container, text="Email", font=("Helvetica", 10))
        email_label.pack(anchor="w")
        self.register_email_var = tk.StringVar()
        email_entry = ttk.Entry(container, textvariable=self.register_email_var, width=35, font=("Helvetica", 10))
        email_entry.pack(pady=(0,10))
        password_label = ttk.Label(container, text="Password", font=("Helvetica", 10))
        password_label.pack(anchor="w")
        self.register_password_var = tk.StringVar()
        password_entry = ttk.Entry(container, textvariable=self.register_password_var, width=35, font=("Helvetica", 10), show="*")
        password_entry.pack(pady=(0,10))
        confirm_label = ttk.Label(container, text="Confirm Password", font=("Helvetica", 10))
        confirm_label.pack(anchor="w")
        self.confirm_password_var = tk.StringVar()
        confirm_entry = ttk.Entry(container, textvariable=self.confirm_password_var, width=35, font=("Helvetica", 10), show="*")
        confirm_entry.pack(pady=(0,20))
        signup_btn = ttk.Button(container, text="Register", command=self.register_user, width=35)
        signup_btn.pack(pady=(0, 10))
        separator = ttk.Label(container, text="or", foreground="#999")
        separator.pack(pady=(10,5))
        google_btn = ttk.Button(container, text="Sign up with Google", command=self.google_sign_in, width=35)
        google_btn.pack(pady=(5, 15))
        bottom_frame = ttk.Frame(container)
        bottom_frame.pack()
        login_label = ttk.Label(bottom_frame, text="Already have an account?", font=("Helvetica", 9))
        login_label.pack(side="left")
        login_link = ttk.Label(bottom_frame, text="Sign in", foreground="blue", cursor="hand2", font=("Helvetica", 9, "underline"))
        login_link.pack(side="left", padx=(5,0))
        login_link.bind("<Button-1>", lambda e: self.show_login_page())
    
    def show_login_page(self):
        self.login_page.tkraise()
    
    def show_register_page(self):
        self.register_page.tkraise()
    
    def login(self):
        username = self.login_username_var.get().strip()
        password = self.login_password_var.get().strip()
        if not username or not password:
            messagebox.showerror("Error", "Please enter both username and password.")
            return
        user = login_user_db(username, password)
        if user:
            messagebox.showinfo("Success", "Login successful!")
            self.current_user = user
            self.load_rules_from_db()
            self.build_rules_tab_ui()  # Rebuild the Rules UI after login
            self.notebook.tab(1, state='normal')
            self.notebook.tab(2, state='normal')
            self.notebook.select(1)
        else:
            messagebox.showerror("Error", "Invalid credentials.")
    
    def register_user(self):
        username = self.register_username_var.get().strip()
        email = self.register_email_var.get().strip()
        password = self.register_password_var.get().strip()
        confirm = self.confirm_password_var.get().strip()
        if not username or not password or not confirm:
            messagebox.showerror("Error", "Please fill all required fields.")
            return
        if password != confirm:
            messagebox.showerror("Error", "Passwords do not match.")
            return
        if os.path.exists(USER_DB):
            with open(USER_DB, "r") as file:
                users = dict(line.strip().split(":") for line in file if ":" in line)
            if username in users:
                messagebox.showerror("Error", "Username already registered.")
                return
        with open(USER_DB, "a") as file:
            file.write(f"{username}:{password}\n")
        messagebox.showinfo("Success", f"Registration successful! You can now log in with username '{username}'.")
        self.show_login_page()
    
    def google_sign_in(self):
        try:
            flow = InstalledAppFlow.from_client_secrets_file(
                'client_secrets.json',
                scopes=['openid',
                        'https://www.googleapis.com/auth/userinfo.profile',
                        'https://www.googleapis.com/auth/userinfo.email']
            )
            credentials = flow.run_local_server(port=0)
            userinfo_endpoint = 'https://www.googleapis.com/oauth2/v1/userinfo'
            params = {'access_token': credentials.token}
            response = requests.get(userinfo_endpoint, params=params)
            if response.status_code == 200:
                user_info = response.json()
                email = user_info.get("email")
                name = user_info.get("name", email)
                if not email:
                    messagebox.showerror("Error", "Google did not return an email address.")
                    return
                if os.path.exists(USER_DB):
                    with open(USER_DB, "r") as file:
                        users = dict(line.strip().split(":") for line in file if ":" in line)
                else:
                    users = {}
                if email in users:
                    messagebox.showinfo("Success", f"Google sign-in successful! Welcome back, {name}.")
                else:
                    with open(USER_DB, "a") as file:
                        file.write(f"{email}:{'google'}\n")
                    messagebox.showinfo("Success", f"Google sign-up successful! Account created for {name}.")
                # For demonstration, set a dummy user dict; in practice, fetch the user record from your DB.
                self.current_user = {"username": email, "id": 1}
                self.load_rules_from_db()
                self.build_rules_tab_ui()
                self.notebook.tab(1, state='normal')
                self.notebook.tab(2, state='normal')
                self.notebook.select(1)
            else:
                messagebox.showerror("Error", "Failed to retrieve user info from Google.")
        except Exception as e:
            messagebox.showerror("Error", f"Google sign-in error: {str(e)}")

if __name__ == "__main__":
    root = tk.Tk()
    app = FileSorterApp(root)
    root.mainloop()
